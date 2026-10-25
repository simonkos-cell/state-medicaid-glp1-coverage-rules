"""
Purpose: build the human-readable Excel summary of state Medicaid weight-loss GLP-1 coverage
         from the two data files, so the workbook is always regenerated, never hand-edited.
Inputs:  data/state_coverage_evidence_sources.csv, data/state_coverage_events.csv
Output:  outputs/state_coverage_summary.xlsx (sheets: by_state, by_state_molecule, events, dimensions, readme)
Date:    2026-10-25
Run:     python3 code/make_summary_workbook.py   (from the repository root)
Assumes: Python 3.10+, openpyxl 3.1.x
"""
import csv, io, os, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = os.path.join(root, "data", "state_coverage_evidence_sources.csv")
evs = os.path.join(root, "data", "state_coverage_events.csv")
out = os.path.join(root, "outputs", "state_coverage_summary.xlsx")
os.makedirs(os.path.dirname(out), exist_ok=True)

rows = list(csv.DictReader(io.StringIO(open(src, encoding="utf-8-sig", newline="").read())))
events = list(csv.DictReader(open(evs, encoding="utf-8-sig", newline="")))
print(f"read {len(rows)} state rows, {len(events)} event rows")
if len(rows) != 50:
    sys.exit(f"expected 50 state rows, got {len(rows)}")

wb = Workbook()
wb.properties.creator = "Kosali Simon"; wb.properties.lastModifiedBy = "Kosali Simon"
head = Font(bold=True); wrap = Alignment(wrap_text=True, vertical="top")
fill = PatternFill("solid", fgColor="DDDDDD")

def sheet(name, header, data, widths=None):
    ws = wb.create_sheet(name)
    ws.append(header)
    for c in ws[1]: c.font = head; c.fill = fill
    for r in data: ws.append(r)
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, 16)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = wrap
    ws.freeze_panes = "B2"
    return ws

# by_state: the compact view most people want
h1 = ["state","coverage_status","wegovy_start","wegovy_end","saxenda_start","saxenda_end","zepbound_start","zepbound_end","prior_authorization","quantity_limits","pa_stringency","w34949_adoption_quarter","khanna_start_quarter","klebanoff_stringency_june2024"]
d1 = [[r["state"],r["coverage_status"],r["wegovy_start_date"],r["wegovy_end_date"],r["saxenda_start_date"],r["saxenda_end_date"],r["zepbound_start_date"],r["zepbound_end_date"],r["prior_authorization"],r["quantity_limits"],r["pa_stringency"],r["w34949_adoption_quarter"],r["khanna_start_quarter"],r["klebanoff_stringency_june2024"]] for r in rows]
sheet("by_state", h1, d1, {"state":18,"wegovy_end":26,"saxenda_end":26,"zepbound_end":26,"pa_stringency":28,"w34949_adoption_quarter":24,"khanna_start_quarter":24,"klebanoff_stringency_june2024":28})

# by_state_molecule: long format, one row per state x molecule with a start date
h2 = ["state","molecule","start_date","end_date","prior_authorization","quantity_limits","pa_stringency","criteria_as_of"]
d2 = []
for r in rows:
    for m in ("wegovy","saxenda","zepbound"):
        s = r[f"{m}_start_date"]
        if s and s != "NA":
            d2.append([r["state"], m.capitalize(), s, r[f"{m}_end_date"], r["prior_authorization"], r["quantity_limits"], r["pa_stringency"], r["criteria_as_of"]])
sheet("by_state_molecule", h2, d2, {"state":18,"end_date":28,"pa_stringency":28})
print(f"by_state_molecule rows: {len(d2)}")

# events
h3 = list(events[0].keys())
sheet("events", h3, [[e[k] for k in h3] for e in events], {"description":50,"source_citation":50,"source_url":40,"notes":40})

# dimensions: the criteria columns for covering states
dims = ["state","pa_stringency","bmi_adult_criteria","bmi_adolescent_criteria","step_therapy","age_minimum","quantity_limit_detail","ffs_mco_scope","criteria_as_of","criteria_source"]
d4 = [[r[k] for k in dims] for r in rows if r["coverage_status"] == "Covered"]
sheet("dimensions", dims, d4, {k:45 for k in dims} | {"state":18,"pa_stringency":24,"age_minimum":16,"criteria_as_of":14})
print(f"dimensions rows (covering states): {len(d4)}")

# readme
ws = wb.active; ws.title = "readme"
lines = [
 "State Medicaid coverage of weight-loss GLP-1s: summary workbook",
 f"Generated {datetime.date.today().isoformat()} by code/make_summary_workbook.py from data/state_coverage_evidence_sources.csv and data/state_coverage_events.csv.",
 "Do not edit this workbook by hand; edit the CSV files and rerun the script.",
 "Sheets: by_state (one row per state), by_state_molecule (one row per state and drug with a start date), events (dated changes), dimensions (PA criteria columns for covering states).",
 "Field definitions: docs/coverage_dimensions_codebook.md. Audit and sources: docs/audit_2026-10.md. Oddities: docs/state_notes.md.",
 "Citation: see CITATION.cff in the repository. Author: Kosali Simon.",
]
for l in lines: ws.append([l])
ws.column_dimensions["A"].width = 140
wb.save(out)
print("wrote", out)
